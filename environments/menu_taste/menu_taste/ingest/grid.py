"""Import spreadsheets where each row is a menu of (title, description) column pairs.

Section headers sit in the same cells as dishes, e.g.

    | PASSED CANAPES | 4 included | Seared Scallop | romesco | SALAD COURSE: Frisée (V/GF) | ...

This walks each row left to right, starts a new section at every header, and turns the
remaining cells into dishes. No model is involved, so nothing leaves the machine.
"""

from __future__ import annotations

import hashlib
import re
from collections import Counter
from dataclasses import replace
from pathlib import Path

from ..schema import (
    Course,
    Dietary,
    EventType,
    Item,
    ItemRole,
    MealPeriod,
    Menu,
    MenuKind,
    Season,
    Section,
    SelectionMode,
    ServiceStyle,
    SourceRef,
)
from .headers import (
    classify_header,
    clean_header,
    looks_like_header,
    split_inline_header,
)
from .redact import redact_menu

DIET_TOKENS = {
    "V": Dietary.vegetarian,
    "VEG": Dietary.vegetarian,
    "VEGETARIAN": Dietary.vegetarian,
    "VG": Dietary.vegan,
    "VEGAN": Dietary.vegan,
    "GF": Dietary.gluten_free,
    "GLUTEN FREE": Dietary.gluten_free,
    "DF": Dietary.dairy_free,
    "DAIRY FREE": Dietary.dairy_free,
    "NF": Dietary.nut_free,
    "NUT FREE": Dietary.nut_free,
}
# Tags with no schema field ("AF" = allergy-friendly); stripped from the name.
IGNORED_DIET_TOKENS = {"AF", "GARLIC", "GARLIC FREE", "SOY FREE"}
_PARENS = re.compile(r"\s*\(([^()]*)\)")
_BREAD = re.compile(
    r"(?<!flat )(?<!flat)\b(bread|breads|baguettes?|focaccia|dinner rolls|epi|sourdough|levain|ciabatta)\b"
    r"(?!\s*(pudding|crumb|salad|crostini|crisp))",
    re.IGNORECASE,
)
# Placeholders and tasting-only extras weren't part of the served menu.
_PLACEHOLDER = re.compile(r"\btbd\b|\btasting only\b|\bnot priced in\b", re.IGNORECASE)
_LEADING_NOTE = re.compile(r"^\s*\([^)]*\)\s*")
_CLIENT_PROVIDED = re.compile(
    r"\b(provided by (?!(the )?(chef|kitchen)\b)|client[- ]provid|(client|couple|host) (to|will) "
    r"(provide|source|bring))",
    re.IGNORECASE,
)
# A special plate for one person ("(BRIDE ONLY)") is an alternate, not part of the menu.
_ONE_PERSON = re.compile(
    r"\s*\((bride|groom|couple|host|dj|officiant)s?\s+only\)", re.IGNORECASE
)
_NOT_GUEST_FOOD = re.compile(
    r"\b(staff|vendor|crew) meals?\b|\binvoice only\b|\bgift tote\b|\braffle\b",
    re.IGNORECASE,
)
_PICK_COUNT = re.compile(
    r"\b(?:choose|choice of|selection of|select|up to|pick)\s*(?:up to\s*)?(\d+)|\b(\d+)\s+included\b",
    re.IGNORECASE,
)
_SIDE_DISH = re.compile(
    r"\b(potato(es)?|fingerlings?|gratin|rice pilaf|pilaf|polenta|fries|mashed|succotash)\b",
    re.IGNORECASE,
)
_PROTEIN = re.compile(
    r"\b(chicken|beef|filet|steak|strip|tri ?tip|short ribs?|pork|lamb|salmon|cod|halibut|fish|"
    r"shrimp|duck|tofu|bass|trout|scallops?|tenderloin|brisket|porchetta|sausage|meatballs?)\b",
    re.IGNORECASE,
)
_MOVABLE = {Course.salad, Course.soup, Course.dessert}
# For sections whose header names no course ("SERVED FAMILY STYLE"), read the dish.
_ITEM_COURSE = [
    (
        re.compile(
            r"\b(potato|potatoes|fries|gratin|rice|polenta|succotash|beans|vegetables|veggies|brussels|carrots|broccolini|corn|mac and cheese|mac & cheese)\b",
            re.IGNORECASE,
        ),
        Course.side,
    ),
    (
        re.compile(r"\b(salad|caesar|greens|slaw|chicories)\b", re.IGNORECASE),
        Course.salad,
    ),
    (re.compile(r"\b(soup|bisque|chowder|gazpacho)\b", re.IGNORECASE), Course.soup),
    (
        re.compile(
            r"\b(board|charcuterie|crudit[eé]s?|cheese (and|&) salumi|snack table|mezze)\b",
            re.IGNORECASE,
        ),
        Course.display,
    ),
    (
        re.compile(
            r"\b(cookies?|brownies?|blondies?|tarts?|cakes?|cupcakes?|pies?|crisp|cobbler|eclairs?|mousse|panna cotta|cheesecakes?|macarons?|truffles|pot de cr[eè]me|crumb bar|curd bar)\b",
            re.IGNORECASE,
        ),
        Course.dessert,
    ),
    (
        re.compile(
            r"\b(pasta|penne|farfalle|rigatoni|tortellini|ravioli|gnocchi|orecchiette|cavatelli|bucatini|lasagna)\b",
            re.IGNORECASE,
        ),
        Course.pasta,
    ),
]
_STATION_ITEM = re.compile(r"\bstation\b", re.IGNORECASE)
_ALLERGY_HEADER = re.compile(
    r"^allergy[- ]friendly (meal|entr[eé]e|option)\b", re.IGNORECASE
)
_CATALOG = re.compile(
    r"\btasting \d", re.IGNORECASE
)  # "choose N" is judged by item count below
_SAMPLE = re.compile(r"\bsample menu\b", re.IGNORECASE)
_JUNK = re.compile(r"^[\W_]*$|^.{0,2}$")


def parse_dietary(name: str) -> tuple[str, list[Dietary]]:
    """'Brownies (GF/VG)' -> ('Brownies', [gluten_free, vegan])."""
    tags: list[Dietary] = []

    def repl(m: re.Match) -> str:
        tokens = [
            t.strip().upper()
            for t in re.split(r"[/,&+]|\band\b", m.group(1))
            if t.strip()
        ]
        known = [DIET_TOKENS[t] for t in tokens if t in DIET_TOKENS]
        ignored = [t for t in tokens if t in IGNORED_DIET_TOKENS]
        if tokens and len(known) + len(ignored) == len(tokens):
            tags.extend(k for k in known if k not in tags)
            return ""
        if known and len(known) >= 2:
            tags.extend(k for k in known if k not in tags)
            return ""
        return m.group(0)

    cleaned = _PARENS.sub(repl, name)

    def star(m: re.Match) -> str:  # "Citrus Salad *GF"
        tag = DIET_TOKENS.get(m.group(1).upper())
        if tag and tag not in tags:
            tags.append(tag)
        return ""

    cleaned = re.sub(r"\s*\*\s*(GF|VG|V|DF|NF)\b", star, cleaned, flags=re.IGNORECASE)
    return re.sub(r"\s{2,}", " ", cleaned).strip(" ,;-"), tags


def load_scrub_terms(path: str | Path | None) -> list[re.Pattern]:
    """One name per line (business, supplier, venue); '#' starts a comment."""
    if not path:
        return []
    terms = []
    for line in Path(path).read_text(encoding="utf-8").splitlines():
        term = line.split("#", 1)[0].strip()
        if term:
            terms.append(re.compile(rf"(?<!\w){re.escape(term)}(?!\w)", re.IGNORECASE))
    return terms


def scrub_names(text: str | None, terms: list[re.Pattern]) -> str | None:
    if not text:
        return text
    for term in terms:
        text = term.sub("", text)
    text = re.sub(r"\s*,\s*(,\s*)+", ", ", text)
    text = re.sub(r"\(\s*\)|\s{2,}", " ", text)
    text = re.sub(r"\b(by|from|with)\s*([,;.]|$)", r"\2", text)
    return text.strip(" ,;:-–") or None


def _event_hints(texts: list[str]) -> dict:
    blob = " ".join(texts).lower()
    hints: dict = {}
    if "wedding" in blob:
        hints["event_type"] = EventType.wedding
    elif "rehearsal" in blob:
        hints["event_type"] = EventType.rehearsal_dinner
    if "christmas" in blob or "holiday" in blob:
        hints.setdefault("event_type", EventType.holiday_party)
        hints["season"] = Season.winter
        if "christmas" in blob:
            hints["month"] = 12
    elif "thanksgiving" in blob:
        hints.update(season=Season.fall, month=11)
    for word, period in (
        ("brunch", MealPeriod.brunch),
        ("breakfast", MealPeriod.breakfast),
        ("lunch", MealPeriod.lunch),
        ("dinner", MealPeriod.dinner),
    ):
        if re.search(rf"\b{word}\b", blob):
            hints["meal_period"] = period
            break
    return hints


def _menu_service(sections: list[Section], drop_off: bool) -> ServiceStyle:
    if drop_off:
        return ServiceStyle.drop_off
    main = {s.service for s in sections if s.course is Course.main and s.service}
    if len(main) == 1:
        return main.pop()
    if len(main) > 1:
        return ServiceStyle.mixed
    styles = Counter(s.service for s in sections if s.service)
    if not styles:
        return ServiceStyle.unknown
    if set(styles) <= {
        ServiceStyle.passed,
        ServiceStyle.display,
        ServiceStyle.stations,
    }:
        return ServiceStyle.reception
    return styles.most_common(1)[0][0]


def _dish_follows(pairs: list[tuple[str, str]], index: int) -> bool:
    for title, desc in pairs[index + 1 :]:
        title = title.strip()
        if not title:
            continue
        head, _ = split_inline_header(title)
        return not (
            head
            or _STATION_ITEM.search(title)
            or looks_like_header(title, desc.strip())
        )
    return False


def parse_row(
    pairs: list[tuple[str, str]], *, menu_id: str, scrub: list[re.Pattern] | None = None
) -> tuple[Menu | None, list[str]]:
    """Turn one spreadsheet row into a menu. Returns (menu or None, warnings)."""
    scrub = scrub or []
    warnings: list[str] = []
    sections: list[Section] = []
    current: Section | None = None
    dropping = False
    role = ItemRole.standard
    kind = MenuKind.event
    drop_off = False
    header_texts: list[str] = []
    notes: list[str] = []
    pick_counts: dict[int, int] = {}

    def start(header: str, description: str = "", inherit_course: bool = False) -> None:
        nonlocal current, dropping, role, kind, drop_off
        info = classify_header(header)
        if (
            inherit_course
            and current is not None
            and info.course in (Course.main, Course.other)
        ):
            info = replace(
                info,
                course=current.course,
                service=info.service or ServiceStyle.stations,
            )
        header_texts.append(header)
        if _CATALOG.search(header) or _CATALOG.search(description):
            kind = MenuKind.catalog
        if info.service is ServiceStyle.drop_off:
            drop_off = True
        if info.action == "attach":
            return
        dropping = info.action == "drop"
        if dropping:
            return
        role = info.role
        current = Section(
            course=info.course,
            heading=clean_header(header) or None,
            service=info.service,
            selection=info.selection,
            items=[],
        )
        sections.append(current)
        if description:
            notes.append(f"{current.heading}: {description}")
            m = _PICK_COUNT.search(description)
            if m:
                pick_counts[id(current)] = int(m.group(1) or m.group(2))

    for index, (raw_title, raw_desc) in enumerate(pairs):
        title, desc = raw_title.strip(), raw_desc.strip()
        if _JUNK.match(title) and _JUNK.match(desc):
            continue
        if _SAMPLE.search(title):
            kind = MenuKind.sample if kind is MenuKind.event else kind
            continue
        if not title:
            warnings.append(f"description without a dish: {desc[:40]!r}")
            continue
        head, rest = split_inline_header(title)
        if head:
            start(head)
            title = rest
        elif _ALLERGY_HEADER.match(title) or looks_like_header(
            title, desc if index else ""
        ):
            start(title, desc)
            continue
        elif _STATION_ITEM.search(title) and _dish_follows(pairs, index):
            # "Ramen Station" followed by "Shoyu Ramen" names the station its
            # dishes are served from; it is a sub-header, not a dish. A station with
            # nothing listed after it ("Gelato Station") is the dish.
            start(title, desc, inherit_course=True)
            continue
        if dropping:
            continue
        if title.startswith("*"):
            notes.append(title.lstrip("* "))  # "*vegan option for the band"
            continue
        if _NOT_GUEST_FOOD.search(title) or _PLACEHOLDER.search(f"{title} {desc}"):
            continue
        if _CLIENT_PROVIDED.search(f"{title} {desc}"):
            header_texts.append(
                title
            )  # "Wedding Cake ... provided by client" is a hint
            continue
        item_role = role
        if _ONE_PERSON.search(title):
            title, item_role = _ONE_PERSON.sub("", title), ItemRole.dietary_alternate
        name, dietary = parse_dietary(_LEADING_NOTE.sub("", title))
        desc, desc_tags = parse_dietary(desc)
        dietary += [t for t in desc_tags if t not in dietary]
        name = scrub_names(name, scrub)
        if not name:
            continue
        item = Item(
            name=name,
            description=scrub_names(desc, scrub) or None,
            dietary=dietary,
            role=item_role,
        )
        if current is None:
            current = Section(course=Course.other, items=[])
            sections.append(current)
        current.items.append(item)

    # "Choose 4" over a list of eight is a proposal; over a list of four, it's the choice.
    for section in sections:
        count = pick_counts.get(id(section))
        if count and len(section.items) > count:
            section.selection = SelectionMode.host_choice
            section.select_count = count
            kind = MenuKind.catalog

    # A dish that names its course ("... Salad", "... Brownies") under an entrée header
    # usually means the export lost a header (boxed lunches list everything in a row).
    for section in list(sections):
        if section.course is not Course.main:
            continue
        moved: dict[Course, list[Item]] = {}
        for item in section.items:
            if _PROTEIN.search(item.name):
                continue
            course = next(
                (c for rx, c in _ITEM_COURSE if c in _MOVABLE and rx.search(item.name)),
                None,
            )
            if course:
                moved.setdefault(course, []).append(item)
        if moved and sum(map(len, moved.values())) < len(section.items):
            gone = [i for items in moved.values() for i in items]
            section.items = [i for i in section.items if i not in gone]
            at = sections.index(section) + 1
            for course, items in moved.items():
                sections.insert(
                    at,
                    section.model_copy(
                        update={
                            "course": course,
                            "heading": None,
                            "items": items,
                            "selection": SelectionMode.all,
                        }
                    ),
                )
                at += 1

    # Sides listed among the entrées ("Potato Pavé") are sides.
    for section in list(sections):
        if section.course is not Course.main:
            continue
        sides = [
            it
            for it in section.items
            if _SIDE_DISH.search(it.name) and not _PROTEIN.search(it.name)
        ]
        if sides and len(sides) < len(section.items):
            section.items = [it for it in section.items if it not in sides]
            sections.insert(
                sections.index(section) + 1,
                section.model_copy(
                    update={
                        "course": Course.side,
                        "heading": None,
                        "items": sides,
                        "selection": SelectionMode.all,
                    }
                ),
            )

    # Sections with no course in their header get one per dish, when the dish says it.
    resolved: list[Section] = []
    for section in sections:
        if section.course is not Course.other:
            resolved.append(section)
            continue
        groups: dict[Course, list[Item]] = {}
        for item in section.items:
            course = next(
                (c for rx, c in _ITEM_COURSE if rx.search(item.name)), Course.other
            )
            groups.setdefault(course, []).append(item)
        for course, items in groups.items():
            resolved.append(
                section.model_copy(update={"course": course, "items": items})
            )
    sections = resolved

    # House bread served with a course is its own course, not a main or a salad.
    breads = []
    for section in sections:
        keep = []
        for item in section.items:
            if section.course is not Course.bread and _BREAD.search(item.name):
                breads.append(item.model_copy(update={"role": ItemRole.standard}))
            else:
                keep.append(item)
        section.items = keep
    sections = [s for s in sections if s.items]
    if breads:
        sections.append(Section(course=Course.bread, items=breads))
    if not sections:
        return None, warnings

    if drop_off:  # "pre-ordered" in a boxed lunch means boxed, not plated
        for section in sections:
            if section.service is ServiceStyle.plated:
                section.service = None

    hints = _event_hints(header_texts)
    menu = Menu(
        menu_id=menu_id,
        kind=kind,
        service=_menu_service(sections, drop_off),
        sections=sections,
        notes=[scrub_names(n, scrub) for n in notes if scrub_names(n, scrub)],
    )
    for key, value in hints.items():
        setattr(menu.event, key, value)
    for section in menu.sections:
        section.heading = scrub_names(section.heading, scrub)
    menu = redact_menu(menu)
    menu.warnings = warnings + menu.warnings
    return menu, warnings


def read_grid(
    path: str | Path, *, pair_columns: int | None = None
) -> list[tuple[int, list[tuple[str, str]]]]:
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.worksheets[0]
    rows = []
    for row_no, values in enumerate(ws.iter_rows(values_only=True), 1):
        cells = ["" if v is None else str(v).strip() for v in values]
        width = pair_columns * 2 if pair_columns else len(cells) + len(cells) % 2
        cells += [""] * (width - len(cells))
        pairs = [(cells[i], cells[i + 1]) for i in range(0, width, 2)]
        if any(a or b for a, b in pairs):
            rows.append((row_no, pairs))
    return rows


def import_grid(
    path: str | Path,
    *,
    scrub_file: str | Path | None = None,
    pair_columns: int | None = None,
) -> tuple[list[Menu], dict]:
    data = Path(path).read_bytes()
    sha = hashlib.sha256(data).hexdigest()
    scrub = load_scrub_terms(scrub_file)
    menus: list[Menu] = []
    seen: dict[tuple, int] = {}
    report = Counter()
    for row_no, pairs in read_grid(path, pair_columns=pair_columns):
        key = tuple((a.lower(), b.lower()) for a, b in pairs if a or b)
        if key in seen:
            report["exact duplicate rows skipped"] += 1
            continue
        seen[key] = row_no
        menu, _ = parse_row(pairs, menu_id=f"{sha[:8]}-r{row_no}", scrub=scrub)
        if menu is None:
            report["rows with no dishes"] += 1
            continue
        menu.source = SourceRef(
            sha256=sha,
            media_type="application/vnd.ms-excel.grid",
            index_in_source=row_no,
        )
        report[f"kind: {menu.kind.value}"] += 1
        menus.append(menu)
    return menus, dict(report)

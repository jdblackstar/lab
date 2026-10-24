"""Turn arbitrary files into model-readable parts (text, images, PDFs).

File types are detected from content, not extensions, because menus arrive renamed,
emailed, zipped, or exported from phones. Supported:

- PDF (sent as-is: the model reads both the text layer and scanned pages)
- Images: PNG, JPEG, GIF, WebP, TIFF (all pages), BMP, HEIC/HEIF (with ``pillow-heif``)
- Word .docx, PowerPoint .pptx (text plus embedded pictures), Excel .xlsx, CSV/TSV
- HTML, plain text, Markdown, RTF
- Email (.eml), including attachments; ZIP archives, recursively
- Legacy Office / OpenDocument (.doc, .xls, .ppt, .odt, ...) via LibreOffice when installed
- Canonical menu JSON / JSONL, passed through without a model call
"""

from __future__ import annotations

import email
import email.policy
import hashlib
import io
import json
import shutil
import subprocess
import tempfile
import zipfile
from dataclasses import dataclass, field
from html.parser import HTMLParser
from pathlib import Path
from xml.etree import ElementTree

MAX_PDF_BYTES = 30 * 1024 * 1024
MAX_IMAGE_EDGE = 2400
MAX_IMAGE_BYTES = 3_500_000
# Archive limits, checked against each member's declared size before it is read.
# Real menu archives are far smaller; these only stop zip bombs.
MAX_ARCHIVE_MEMBER_BYTES = 100 * 1024 * 1024
MAX_ARCHIVE_TOTAL_BYTES = 500 * 1024 * 1024  # across all nested archives in one file
MAX_ARCHIVE_MEMBERS = 5000
MAX_ARCHIVE_DEPTH = 3  # a zip inside an email inside a zip is depth 3
# Office files with less text than this are probably a pasted picture of a menu.
SPARSE_TEXT_CHARS = 200

NATIVE_IMAGE_TYPES = {"image/png", "image/jpeg", "image/gif", "image/webp"}


@dataclass
class Part:
    kind: str  # "text" | "image" | "pdf"
    data: str | bytes
    media_type: str
    label: str


@dataclass
class SourceDoc:
    """One logical document. ``path`` may name a member inside an archive or email."""

    path: str
    sha256: str
    media_type: str
    parts: list[Part] = field(default_factory=list)
    menus_json: list[dict] | None = None  # already-canonical records
    notes: list[str] = field(default_factory=list)
    error: str | None = None


class UnsupportedSource(Exception):
    pass


def sniff(data: bytes, name: str) -> str:
    """Return a media type for ``data``, checking magic bytes before the extension."""
    ext = Path(name).suffix.lower()
    head = data[:16]
    if head.startswith(b"%PDF"):
        return "application/pdf"
    if head.startswith(b"\x89PNG"):
        return "image/png"
    if head.startswith(b"\xff\xd8\xff"):
        return "image/jpeg"
    if head.startswith((b"GIF87a", b"GIF89a")):
        return "image/gif"
    if head[:4] == b"RIFF" and head[8:12] == b"WEBP":
        return "image/webp"
    if head.startswith((b"II*\x00", b"MM\x00*")):
        return "image/tiff"
    if head.startswith(b"BM"):
        return "image/bmp"
    if head[4:8] == b"ftyp" and head[8:12] in (
        b"heic",
        b"heix",
        b"mif1",
        b"msf1",
        b"heif",
    ):
        return "image/heic"
    if head.startswith(b"PK\x03\x04"):
        try:
            names = set(zipfile.ZipFile(io.BytesIO(data)).namelist())
        except zipfile.BadZipFile:
            return "application/octet-stream"
        if "word/document.xml" in names:
            return "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        if "xl/workbook.xml" in names:
            return "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        if "ppt/presentation.xml" in names:
            return "application/vnd.openxmlformats-officedocument.presentationml.presentation"
        if "mimetype" in names and "content.xml" in names:
            return "application/vnd.oasis.opendocument"
        if any(n.startswith("Index/") for n in names) and ext in (
            ".pages",
            ".numbers",
            ".key",
        ):
            return "application/x-iwork"
        return "application/zip"
    if head.startswith(b"\xd0\xcf\x11\xe0"):
        return (
            "application/vnd.ms-outlook"
            if ext == ".msg"
            else "application/x-ole-office"
        )
    if head.startswith(b"{\\rtf"):
        return "application/rtf"
    text = _decode(data[:4096])
    stripped = text.lstrip().lower()
    if ext in (".json", ".jsonl") or stripped.startswith(("{", "[")):
        return "application/json"
    if ext in (".html", ".htm") or stripped.startswith(("<!doctype html", "<html")):
        return "text/html"
    if ext == ".eml" or (
        any(
            stripped.startswith(h)
            for h in ("from:", "received:", "mime-version:", "return-path:")
        )
    ):
        return "message/rfc822"
    if ext in (".csv", ".tsv"):
        return "text/csv"
    return "text/plain"


def _decode(data: bytes) -> str:
    for encoding in ("utf-8-sig", "cp1252"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("latin-1")


def load_path(path: str | Path) -> list[SourceDoc]:
    path = Path(path)
    try:
        data = path.read_bytes()
    except OSError as e:
        # Unreadable, or gone since the walk: record it as skipped, don't abort.
        return [SourceDoc(path=str(path), sha256="", media_type="", error=str(e))]
    return load_bytes(data, str(path))


@dataclass
class _ArchiveBudget:
    """What is left to expand from one top-level file, shared by its nested archives."""

    bytes_left: int
    members_left: int


def load_bytes(
    data: bytes, name: str, *, depth: int = 0, budget: _ArchiveBudget | None = None
) -> list[SourceDoc]:
    media_type = sniff(data, name)
    sha = hashlib.sha256(data).hexdigest()
    doc = SourceDoc(path=name, sha256=sha, media_type=media_type)
    budget = budget or _ArchiveBudget(MAX_ARCHIVE_TOTAL_BYTES, MAX_ARCHIVE_MEMBERS)
    try:
        if media_type in ("application/zip", "message/rfc822"):
            if depth >= MAX_ARCHIVE_DEPTH:
                raise UnsupportedSource(
                    f"nested more than {MAX_ARCHIVE_DEPTH} archives/emails deep"
                )
            if media_type == "application/zip":
                return _load_zip(data, name, depth, budget)
            return _load_email(data, name, doc, depth, budget)
        _fill(doc, data, media_type, name)
    except UnsupportedSource as e:
        doc.error = str(e)
    except Exception as e:  # noqa: BLE001 - a corrupt file shouldn't stop a batch
        doc.error = f"{type(e).__name__}: {e}"
    return [doc]


def _fill(doc: SourceDoc, data: bytes, media_type: str, name: str) -> None:
    label = Path(name).name
    if media_type == "application/pdf":
        if len(data) > MAX_PDF_BYTES:
            raise UnsupportedSource(
                f"PDF is {len(data) / 1e6:.0f} MB; split it under {MAX_PDF_BYTES / 1e6:.0f} MB"
            )
        doc.parts.append(Part("pdf", data, media_type, label))
    elif media_type.startswith("image/"):
        doc.parts += image_parts(data, media_type, label)
    elif media_type.endswith("wordprocessingml.document"):
        _office_zip(doc, data, label, _docx_text)
    elif media_type.endswith("presentationml.presentation"):
        _office_zip(doc, data, label, _pptx_text)
    elif media_type.endswith("spreadsheetml.sheet"):
        doc.parts.append(Part("text", _xlsx_text(data), "text/plain", label))
    elif media_type in (
        "application/x-ole-office",
        "application/vnd.oasis.opendocument",
    ):
        pdf = convert_with_libreoffice(data, name)
        doc.notes.append("converted to PDF with LibreOffice")
        doc.parts.append(Part("pdf", pdf, "application/pdf", label))
    elif media_type == "application/vnd.ms-outlook":
        raise UnsupportedSource(
            "Outlook .msg: save the email as .eml or export the attachment"
        )
    elif media_type == "application/x-iwork":
        raise UnsupportedSource("Apple Pages/Numbers/Keynote: export to PDF first")
    elif media_type == "application/rtf":
        from striprtf.striprtf import rtf_to_text

        doc.parts.append(Part("text", rtf_to_text(_decode(data)), "text/plain", label))
    elif media_type == "text/html":
        doc.parts.append(Part("text", html_to_text(_decode(data)), "text/plain", label))
    elif media_type == "application/json":
        try:
            doc.menus_json = _json_records(_decode(data))
        except json.JSONDecodeError:  # text that merely starts with a bracket
            doc.media_type = "text/plain"
            doc.parts.append(Part("text", _decode(data), "text/plain", label))
    elif media_type.startswith("text/"):
        doc.parts.append(Part("text", _decode(data), "text/plain", label))
    else:
        raise UnsupportedSource(f"unrecognized file type ({media_type})")
    if not doc.parts and doc.menus_json is None:
        raise UnsupportedSource("no readable content")


def image_parts(data: bytes, media_type: str, label: str) -> list[Part]:
    """Normalize images: apply phone rotation, split TIFF pages, keep under API limits."""
    from PIL import Image, ImageOps

    if media_type == "image/heic":
        try:
            from pillow_heif import register_heif_opener
        except ImportError as e:
            raise UnsupportedSource("HEIC photo: install pillow-heif") from e
        register_heif_opener()

    if media_type in NATIVE_IMAGE_TYPES and len(data) <= MAX_IMAGE_BYTES:
        with Image.open(io.BytesIO(data)) as img:
            needs_rotation = img.getexif().get(0x0112, 1) != 1
            fits = max(img.size) <= MAX_IMAGE_EDGE
        if fits and not needs_rotation:
            return [Part("image", data, media_type, label)]

    parts = []
    with Image.open(io.BytesIO(data)) as img:
        frames = getattr(img, "n_frames", 1) if media_type == "image/tiff" else 1
        for frame in range(frames):
            img.seek(frame)
            page = ImageOps.exif_transpose(img.copy())
            page.thumbnail((MAX_IMAGE_EDGE, MAX_IMAGE_EDGE))
            if page.mode not in ("RGB", "L"):
                page = page.convert("RGB")
            buf = io.BytesIO()
            page.save(buf, format="JPEG", quality=88)
            suffix = f" (page {frame + 1})" if frames > 1 else ""
            parts.append(Part("image", buf.getvalue(), "image/jpeg", label + suffix))
    return parts


_W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
_A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"


def _docx_text(zf: zipfile.ZipFile) -> str:
    root = ElementTree.fromstring(zf.read("word/document.xml"))
    body = root.find(f"{_W}body")
    lines: list[str] = []

    def para(p) -> str:
        return "".join(t.text or "" for t in p.iter(f"{_W}t")).strip()

    for child in body if body is not None else []:
        if child.tag == f"{_W}p":
            lines.append(para(child))
        elif child.tag == f"{_W}tbl":
            for row in child.iter(f"{_W}tr"):
                cells = [
                    " ".join(filter(None, (para(p) for p in tc.iter(f"{_W}p"))))
                    for tc in row.iter(f"{_W}tc")
                ]
                lines.append(" | ".join(cells))
    # Text boxes and headers often hold the menu title or a whole designed layout.
    for extra in sorted(n for n in zf.namelist() if n.startswith("word/header")):
        header = " ".join(
            t.text or "" for t in ElementTree.fromstring(zf.read(extra)).iter(f"{_W}t")
        ).strip()
        if header:
            lines.insert(0, header)
    return "\n".join(lines).strip()


def _pptx_text(zf: zipfile.ZipFile) -> str:
    def slide_no(n: str) -> int:
        return int("".join(ch for ch in Path(n).stem if ch.isdigit()) or 0)

    slides = sorted(
        (
            n
            for n in zf.namelist()
            if n.startswith("ppt/slides/slide") and n.endswith(".xml")
        ),
        key=slide_no,
    )
    out = []
    for n in slides:
        root = ElementTree.fromstring(zf.read(n))
        paras = [
            "".join(t.text or "" for t in p.iter(f"{_A}t")).strip()
            for p in root.iter(f"{_A}p")
        ]
        out.append(f"[slide {slide_no(n)}]\n" + "\n".join(p for p in paras if p))
    return "\n\n".join(out).strip()


def _office_zip(doc: SourceDoc, data: bytes, label: str, text_fn) -> None:
    zf = zipfile.ZipFile(io.BytesIO(data))
    text = text_fn(zf)
    if text:
        doc.parts.append(Part("text", text, "text/plain", label))
    if len(text) < SPARSE_TEXT_CHARS:
        read = 0
        for n in sorted(n for n in zf.namelist() if "/media/" in n):
            blob = zf.read(n)
            kind = sniff(blob, n)
            if not kind.startswith("image/"):
                continue  # EMF/WMF vector art and similar
            doc.parts += image_parts(blob, kind, f"{label}: {Path(n).name}")
            read += 1
        if read:
            doc.notes.append(f"little text, so read {read} embedded image(s)")


def _xlsx_text(data: bytes) -> str:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    out = []
    for ws in wb.worksheets:
        rows = []
        for row in ws.iter_rows(values_only=True):
            cells = ["" if v is None else str(v).strip() for v in row]
            while cells and not cells[-1]:
                cells.pop()
            if any(cells):
                rows.append("\t".join(cells))
        if rows:
            out.append(f"[sheet {ws.title}]\n" + "\n".join(rows))
    return "\n\n".join(out)


class _TextExtractor(HTMLParser):
    BLOCK = frozenset(
        {
            "p",
            "div",
            "br",
            "li",
            "tr",
            "h1",
            "h2",
            "h3",
            "h4",
            "h5",
            "h6",
            "section",
            "table",
        }
    )

    def __init__(self):
        super().__init__()
        self.out: list[str] = []
        self.skip = 0

    def handle_starttag(self, tag, attrs):
        if tag in ("script", "style", "noscript"):
            self.skip += 1
        elif tag in self.BLOCK:
            self.out.append("\n")
        elif tag in ("td", "th"):
            self.out.append(" | ")

    def handle_endtag(self, tag):
        if tag in ("script", "style", "noscript"):
            self.skip = max(0, self.skip - 1)
        elif tag in self.BLOCK:
            self.out.append("\n")

    def handle_data(self, data):
        if not self.skip:
            self.out.append(data)


def html_to_text(html: str) -> str:
    parser = _TextExtractor()
    parser.feed(html)
    lines = (" ".join(line.split()) for line in "".join(parser.out).splitlines())
    return "\n".join(line for line in lines if line)


def _json_records(text: str) -> list[dict]:
    text = text.strip()
    if text.startswith("["):
        return json.loads(text)
    if "\n" in text and text.startswith("{") and not text.endswith("]"):
        try:
            return [json.loads(line) for line in text.splitlines() if line.strip()]
        except json.JSONDecodeError:
            pass
    return [json.loads(text)]


def _load_zip(
    data: bytes, name: str, depth: int, budget: _ArchiveBudget
) -> list[SourceDoc]:
    docs = []
    with zipfile.ZipFile(io.BytesIO(data)) as zf:
        for info in zf.infolist():
            member = Path(info.filename)
            if (
                info.is_dir()
                or member.name.startswith(".")
                or "__MACOSX" in member.parts
            ):
                continue
            member_name = f"{name}!{info.filename}"
            # zipfile stops at the declared size, so checking it bounds the read.
            budget.members_left -= 1
            budget.bytes_left -= info.file_size
            if budget.members_left < 0 or budget.bytes_left < 0:
                docs.append(
                    _skipped(name, "archive expands past the size or file-count limit")
                )
                break
            if info.file_size > MAX_ARCHIVE_MEMBER_BYTES:
                docs.append(
                    _skipped(member_name, f"{info.file_size / 1e6:.0f} MB uncompressed")
                )
                continue
            try:
                member_data = zf.read(info)
            except Exception as e:  # noqa: BLE001 - an encrypted or corrupt member
                docs.append(_skipped(member_name, f"{type(e).__name__}: {e}"))
                continue
            docs += load_bytes(member_data, member_name, depth=depth + 1, budget=budget)
    return docs


def _skipped(name: str, reason: str) -> SourceDoc:
    return SourceDoc(path=name, sha256="", media_type="", error=f"skipped: {reason}")


def _load_email(
    data: bytes, name: str, doc: SourceDoc, depth: int, budget: _ArchiveBudget
) -> list[SourceDoc]:
    msg = email.message_from_bytes(data, policy=email.policy.default)
    docs = []
    # Mail clients often send a stub plain-text part next to the real HTML body, so
    # keep whichever version has more to say.
    bodies = []
    for kind in ("plain", "html"):
        part = msg.get_body(preferencelist=(kind,))
        if part is not None:
            content = part.get_content()
            bodies.append(html_to_text(content) if kind == "html" else content.strip())
    text = max(bodies, key=len, default="")
    if text:
        doc.parts.append(Part("text", text, "text/plain", f"{Path(name).name} (body)"))
        docs.append(doc)
    for att in msg.iter_attachments():
        payload = att.get_payload(decode=True)
        if payload:
            docs += load_bytes(
                payload,
                f"{name}!{att.get_filename() or 'attachment'}",
                depth=depth + 1,
                budget=budget,
            )
    return docs


def convert_with_libreoffice(data: bytes, name: str) -> bytes:
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        raise UnsupportedSource(
            "legacy Office/OpenDocument file: install LibreOffice or re-save as PDF/.docx"
        )
    with tempfile.TemporaryDirectory() as tmp:
        src = Path(tmp) / Path(name.split("!")[-1]).name
        src.write_bytes(data)
        subprocess.run(
            [soffice, "--headless", "--convert-to", "pdf", "--outdir", tmp, str(src)],
            check=True,
            capture_output=True,
            timeout=180,
        )
        return (Path(tmp) / (src.stem + ".pdf")).read_bytes()

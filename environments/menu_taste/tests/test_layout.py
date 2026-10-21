from menu_taste.ingest.layout import describe_workbook
from openpyxl import Workbook
from openpyxl.styles import Font


def test_layout_hides_every_cell_but_keeps_structure(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.append(["The Goldberg Wedding", "June 14"])
    ws.append(["Course", "Dish", "Price"])
    ws.append(["Main", "Braised short rib, celery root", 42])
    ws["A3"].font = Font(bold=True)
    ws["C3"].number_format = '"$"#,##0.00'
    path = tmp_path / "menus.xlsx"
    wb.save(path)

    out = describe_workbook(path)
    for secret in ("Goldberg", "June", "short rib", "Course", "42"):
        assert secret not in out
    assert "T4b" in out and "T30" in out and "    $" in out

    with_headers = describe_workbook(path, show_headers=True)
    assert "header text: Course | Dish | Price" in with_headers
    assert "Goldberg" not in with_headers and "short rib" not in with_headers
